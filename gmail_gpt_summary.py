import os
import pickle
import base64
import openai
import smtplib
from bs4 import BeautifulSoup
from email.mime.text import MIMEText
from email.header import decode_header
from email.utils import parsedate_to_datetime
from google.auth.transport.requests import Request
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from docx import Document
import openpyxl
import PyPDF2
from dotenv import load_dotenv

# ✅ 환경 변수 불러오기
load_dotenv()

# ✅ 사용자 설정
GMAIL_USER = "th.hwang@koreasmt.co.kr"
TO_EMAIL = "th.hwang@koreasmt.co.kr"
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
openai.api_key = OPENAI_API_KEY

# ✅ 구글 API 권한 설정
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]

def authenticate_gmail():
    creds = None
    if os.path.exists("token.pickle"):
        with open("token.pickle", "rb") as token:
            creds = pickle.load(token)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file("gmail_credentials.json", SCOPES)
            creds = flow.run_local_server(port=0)
        with open("token.pickle", "wb") as token:
            pickle.dump(creds, token)
    return build("gmail", "v1", credentials=creds)

def clean(text):
    return "".join(c for c in text if c.isalnum() or c in (" ", "-", "_")).strip()

def extract_text_from_attachment(filename, data):
    filepath = os.path.join(os.getenv("TEMP"), filename)
    with open(filepath, "wb") as f:
        f.write(data)
    ext = filename.lower().split(".")[-1]
    try:
        if ext == "txt":
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()[:1500]
        elif ext == "pdf":
            with open(filepath, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                return "".join(p.extract_text() for p in reader.pages[:3])[:1500]
        elif ext == "docx":
            doc = Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs)[:1500]
        elif ext == "xlsx":
            wb = openpyxl.load_workbook(filepath)
            sheet = wb.active
            rows = []
            for row in sheet.iter_rows(values_only=True, max_row=10):
                rows.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
            return "\n".join(rows)[:1500]
    except:
        return "[첨부파일 읽기 실패]"
    return ""

def fetch_and_summarize():
    service = authenticate_gmail()
    results = service.users().messages().list(userId="me", labelIds=["INBOX"], q="is:unread", maxResults=10).execute()
    messages = results.get("messages", [])
    summaries = []

    for msg in messages:
        msg_data = service.users().messages().get(userId="me", id=msg['id'], format="full").execute()
        headers = msg_data["payload"].get("headers", [])
        subject = sender = date = "(알 수 없음)"

        for h in headers:
            if h["name"] == "Subject":
                subject = decode_header(h["value"])[0][0]
                if isinstance(subject, bytes): subject = subject.decode()
            if h["name"] == "From":
                sender = h["value"]
            if h["name"] == "Date":
                date = parsedate_to_datetime(h["value"]).strftime("%Y-%m-%d %H:%M")

        if GMAIL_USER not in sender:
            continue

        body = ""
        parts = msg_data["payload"].get("parts", [])
        for part in parts:
            if part["mimeType"] == "text/plain":
                data = base64.urlsafe_b64decode(part["body"]["data"]).decode("utf-8", errors="ignore")
                body = data[:2000]
            elif part["mimeType"] == "text/html":
                html = base64.urlsafe_b64decode(part["body"]["data"]).decode("utf-8", errors="ignore")
                soup = BeautifulSoup(html, "html.parser")
                body = soup.get_text()[:2000]

        attach_texts = []
        for part in parts:
            filename = part.get("filename")
            if filename and part.get("body", {}).get("attachmentId"):
                attach_id = part["body"]["attachmentId"]
                attach = service.users().messages().attachments().get(userId="me", messageId=msg['id'], id=attach_id).execute()
                file_data = base64.urlsafe_b64decode(attach["data"])
                text = extract_text_from_attachment(filename, file_data)
                attach_texts.append(f"[첨부파일 요약: {filename}]\n{text}")

        full_content = body + "\n\n" + "\n".join(attach_texts)

        try:
            response = openai.ChatCompletion.create(
                model="gpt-3.5-turbo",
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "너는 전문 이메일 요약 비서야. 본문과 첨부파일을 분석해서 아래 항목을 빠짐없이 정리해줘.\n"
                            "[요약 포맷]\n"
                            "① 메일 목적 요약 (요청/지시/보고 등 구분)\n"
                            "② 업무 요청 또는 확인 사항 (구체적으로)\n"
                            "③ 일정 및 마감 기한 여부\n"
                            "④ 중요한 수치나 문장 강조\n"
                            "⑤ 관련자 또는 담당자 명시\n"
                            "⑥ [첨부파일 요약:] 파일명별 주요 내용 정리\n"
                            "⑦ 수신자가 반드시 확인해야 할 사항은 맨 아래에 강조해서 요약\n"
                            "항목별로 명확히 정리하고 불필요한 인사말이나 반복은 제거해."
                        )
                    },
                    {"role": "user", "content": full_content}
                ]
            )
            summary = response["choices"][0]["message"]["content"]
            title = f"[GPT 요약] {subject} - {date}"
            summaries.append((title, summary))
        except Exception as e:
            summaries.append((f"[GPT 요약 실패] {subject}", str(e)))

    return summaries

def send_summary(summaries):
    for title, body in summaries:
        msg = MIMEText(body, _charset="utf-8")
        msg["Subject"] = title
        msg["From"] = GMAIL_USER
        msg["To"] = TO_EMAIL
        try:
            server = smtplib.SMTP_SSL("smtp.gmail.com", 465)
            server.login(GMAIL_USER, os.getenv("GMAIL_APP_PASSWORD"))
            server.send_message(msg)
            server.quit()
            print(f"✅ 메일 전송 완료: {title}")
        except Exception as e:
            print(f"❌ 메일 전송 실패: {title} / {e}")

if __name__ == "__main__":
    result = fetch_and_summarize()
    send_summary(result)

